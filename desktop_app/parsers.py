"""Source-preserving adapters for the legacy PDF/supplement operations.

No legacy main, macros, formulas or document instructions are executed.
"""
from __future__ import annotations

import csv
import ctypes
import hashlib
import io
import math
import posixpath
import re
import zipfile
from datetime import date,datetime
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

import pypdfium2 as pdfium
from PIL import Image, ImageOps, ImageDraw, ImageFont

from .reports import PDF_LOCK
from .storage import sha256

MAX_ROWS, MAX_COLUMNS, MAX_PAGES = 3000, 80, 100
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_TYPES = {".txt", ".md", ".json", ".xml", ".html", ".htm"}
TABLE_TYPES = {".xlsx", ".xlsm", ".xls", ".csv", ".tsv"}


def numeric(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        v = float(str(value).strip().replace(",", ""))
        return v if math.isfinite(v) else None
    except (ValueError, TypeError):
        return None


def decode_text(raw):
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def safe_archive(path, max_bytes=80 * 1024 * 1024):
    with zipfile.ZipFile(path) as z:
        members = z.infolist()
        if len(members) > 1200 or sum(v.file_size for v in members) > max_bytes:
            raise ValueError("压缩内容超过解析限额")
        if any(v.file_size > 4 * 1024 * 1024 and v.file_size / max(v.compress_size, 1) > 250 for v in members):
            raise ValueError("压缩率超过解析限额")


def read_tables(path):
    """Return bounded tables, displayed values and exact worksheet coordinates."""
    path = Path(path)
    suffix = path.suffix.lower()
    tables, notes = [], []
    if path.stat().st_size > 24 * 1024 * 1024:
        raise ValueError("表格超过24MB解析限额")
    if suffix in {".csv", ".tsv"}:
        text = decode_text(path.read_bytes())
        reader = csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",")
        rows = []
        for index, row in enumerate(reader):
            if index >= MAX_ROWS:
                notes.append(f"分隔表仅处理前{MAX_ROWS}行")
                break
            if len(row) > MAX_COLUMNS:
                notes.append(f"分隔表仅处理前{MAX_COLUMNS}列")
            rows.append(row[:MAX_COLUMNS])
        return [{"sheet": "CSV", "rows": rows, "hidden": False, "formulas": [], "hidden_ranges": []}], list(dict.fromkeys(notes))
    if suffix == ".xls" and not zipfile.is_zipfile(path):
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets()[:24]:
                rows = [[sheet.cell_value(r, c) for c in range(min(sheet.ncols, MAX_COLUMNS))] for r in range(min(sheet.nrows, MAX_ROWS))]
                tables.append({"sheet": sheet.name, "rows": rows, "hidden": bool(getattr(sheet, "visibility", 0)), "formulas": [], "hidden_ranges": []})
                if sheet.nrows > MAX_ROWS or sheet.ncols > MAX_COLUMNS:
                    notes.append(f"{sheet.name}仅处理前{MAX_ROWS}行/{MAX_COLUMNS}列")
            notes.append("旧版XLS的公式、日期格式及隐藏行列未完整解析")
        finally:
            book.release_resources()
        return tables, notes
    import openpyxl
    safe_archive(path)
    # BytesIO also handles OOXML files mislabeled as .xls.
    raw = path.read_bytes()
    formulas = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True, keep_links=False)
    try:
        for ws in formulas.worksheets[:24]:
            nr, nc = min(ws.max_row, MAX_ROWS), min(ws.max_column, MAX_COLUMNS)
            values = list(cached[ws.title].iter_rows(max_row=nr, max_col=nc, values_only=True))
            rows, formula_cells, date_cells = [], [], []
            for ri, row in enumerate(ws.iter_rows(max_row=nr, max_col=nc)):
                out = []
                for ci, cell in enumerate(row):
                    value = cell.value
                    if isinstance(value,(date,datetime)):
                        date_cells.append({"row":ri+1,"column":ci+1,"value":str(value)})
                    if cell.data_type == "f":
                        value = values[ri][ci] if ri < len(values) else None
                        formula_cells.append({"cell": cell.coordinate, "expression": str(cell.value)[:200], "cached": value is not None})
                    if isinstance(value, (float, int)) and not isinstance(value, bool) and re.fullmatch(r"0\.0+", cell.number_format or ""):
                        value = f"{value:.{len(cell.number_format)-2}f}"
                    elif value is not None and not isinstance(value, (str, int, float, bool)):
                        value = str(value)
                    out.append(value if value is not None else "")
                rows.append(out)
            hidden = [f"row:{i}" for i,d in ws.row_dimensions.items() if d.hidden] + [f"column:{i}" for i,d in ws.column_dimensions.items() if d.hidden]
            tables.append({"sheet": ws.title, "rows": rows, "hidden": ws.sheet_state != "visible", "hidden_ranges": hidden, "formulas": formula_cells, "date_cells":date_cells})
            if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLUMNS:
                notes.append(f"{ws.title}仅处理前{MAX_ROWS}行/{MAX_COLUMNS}列")
        if len(formulas.worksheets) > 24:
            notes.append("工作簿仅处理前24张工作表")
    finally:
        formulas.close()
        cached.close()
    return tables, notes


def pdf_box(page, bounds, width, height):
    points = []
    l, b, r, t = bounds
    for x,y in [(l,b), (l,t), (r,b), (r,t)]:
        dx, dy = ctypes.c_int(), ctypes.c_int()
        if not pdfium.raw.FPDF_PageToDevice(page, 0, 0, width, height, 0, x, y, ctypes.byref(dx), ctypes.byref(dy)):
            raise ValueError("PDF坐标映射失败")
        points.append((dx.value / width, dy.value / height))
    return [max(0., min(p[0] for p in points)), max(0., min(p[1] for p in points)), min(1., max(p[0] for p in points)), min(1., max(p[1] for p in points))]


class Parser:
    def __init__(self, store, progress=lambda *_: None, cancelled=lambda: False):
        self.store, self.progress, self.cancelled = store, progress, cancelled
        self.segments, self.tables, self.images, self.coverage, self.notes = [], [], [], [], []
        self.derived_ids=[]

    def image(self, image, artifact, page, bbox, kind):
        if len(self.images) >= 240:
            if "图像索引限240项，额外图像未进入匹配" not in self.notes:
                self.notes.append("图像索引限240项，额外图像未进入匹配")
            return
        if min(image.size) < 48:
            return
        image = image.convert("RGB")
        image.thumbnail((1250, 1250))
        key = hashlib.sha256((artifact["id"] + str(page) + str(bbox) + kind).encode()).hexdigest()[:24]
        target = self.store.resolve(f"papers/{artifact['paper_id']}/parsed/images/{key}.png")
        target.parent.mkdir(parents=True, exist_ok=True)
        image.save(target)
        self.images.append({"id": key, "artifact_id": artifact["id"], "page": page, "bbox": bbox, "kind": kind, "path": str(target), "source_sha256": artifact["sha256"]})

    def parse(self, artifacts):
        for artifact in artifacts:
            if self.cancelled():
                raise InterruptedError()
            self.one(artifact)
        return self

    def one(self, artifact, depth=0):
        path = self.store.resolve(artifact["relative_path"])
        suffix = path.suffix.lower()
        record = {"artifact_id": artifact["id"], "name": artifact["name"], "status": "processed"}
        self.coverage.append(record)
        self.progress("parsing", "正在解析：" + artifact["name"])
        try:
            if sha256(path) != artifact["sha256"]:
                raise ValueError("文件哈希已变化")
            if suffix == ".pdf":
                self.pdf(path, artifact, record)
            elif suffix in TABLE_TYPES:
                tables, notes = read_tables(path)
                for table in tables:
                    table["artifact_id"] = artifact["id"]
                    table["source_sha256"] = artifact["sha256"]
                    self.tables.append(table)
                    for start in range(0, len(table["rows"]), 50):
                        text = "\n".join(f"R{r+start+1}: " + " | ".join(str(v)[:80] for v in row) for r,row in enumerate(table["rows"][start:start+50]))
                        self.segments.append({"artifact_id": artifact["id"], "sheet": table["sheet"], "row_start": start+1, "text": text})
                self.notes.extend(artifact["name"] + "：" + n for n in notes)
                record.update(tables=len(tables), status="partial" if notes else "processed")
            elif suffix in IMAGE_TYPES:
                with Image.open(path) as source:
                    count = getattr(source, "n_frames", 1)
                    for frame in range(min(count, 24)):
                        source.seek(frame)
                        self.image(ImageOps.exif_transpose(source).copy(), artifact, frame+1, [0,0,1,1], "raster")
                    record.update(frames_processed=min(count,24), frames_total=count)
                    if count > 24:
                        record["status"] = "partial"
                        self.notes.append(artifact["name"] + "：仅解析前24帧")
            elif suffix == ".zip" and depth == 0:
                safe_archive(path)
                with zipfile.ZipFile(path) as z:
                    members = [v for v in z.infolist() if not v.is_dir()]
                    if len(members) > 200:
                        raise ValueError("ZIP超过200个文件")
                    names = set()
                    for member in members:
                        name = PurePosixPath(member.filename.replace("\\", "/"))
                        if name.is_absolute() or ".." in name.parts or ":" in str(name) or str(name) in names:
                            raise ValueError("ZIP成员路径不合法或重复")
                        names.add(str(name))
                        allowed = {".pdf", ".docx"} | TABLE_TYPES | IMAGE_TYPES | TEXT_TYPES
                        if name.suffix.lower() not in allowed:
                            self.notes.append(f"{artifact['name']} / {name}：不支持该成员，未分析")
                            record["status"] = "partial"
                            continue
                        child = self.store.register_child(artifact, str(name), z.read(member))
                        self.derived_ids.append(child['id'])
                        self.one(child, depth+1)
                    record["members"] = len(members)
            elif suffix == ".docx":
                safe_archive(path)
                with zipfile.ZipFile(path) as z:
                    xml = ElementTree.fromstring(z.read("word/document.xml"))
                    rels={}
                    if 'word/_rels/document.xml.rels' in z.namelist():
                        for r in ElementTree.fromstring(z.read('word/_rels/document.xml.rels')):
                            if r.get('TargetMode')!='External' and r.get('Type','').endswith('/image'):
                                rels[r.get('Id')]=posixpath.normpath(posixpath.join('word',r.get('Target','')))
                    seen=set();extracted=0
                    paragraphs=[p for p in xml.iter() if p.tag.endswith('}p')]
                    for index, paragraph in enumerate(paragraphs,1):
                        text=''.join(t.text or '' for t in paragraph.iter() if t.tag.endswith('}t'))
                        if text.strip():self.segments.append({'artifact_id':artifact['id'],'paragraph':index,'text':text})
                        for el in paragraph.iter():
                            rid=el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') or el.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                            name=rels.get(rid)
                            if not name or name in seen:continue
                            seen.add(name)
                            if not name.startswith('word/media/') or '..' in PurePosixPath(name).parts or Path(name).suffix.lower() not in IMAGE_TYPES:
                                self.notes.append(artifact['name']+'：部分嵌入图片格式未支持');continue
                            if extracted>=80:self.notes.append(artifact['name']+'：嵌入图仅处理前80项');break
                            child=self.store.register_child(artifact,name,z.read(name))
                            self.derived_ids.append(child['id'])
                            self.segments.append({'artifact_id':child['id'],'source_type':'derived_provenance','parent_artifact_id':artifact['id'],'parent_paragraph':index,'text':f"来源附件 {artifact['name']}，段落{index}嵌入原图 {name}。邻近文字：{text}"})
                            self.one(child,depth+1);extracted+=1
                    record.update(embedded_images=extracted,paragraphs=len(paragraphs))
                self.notes.append(artifact["name"] + "：已读取文字及正文引用的支持格式嵌入图；Word表格未结构化，浮动图文的精确对应仍需人工核对")
                record["status"] = "partial"
            elif suffix in TEXT_TYPES:
                if path.stat().st_size > 8*1024*1024:
                    raise ValueError("文本超过8MB解析限额")
                text = decode_text(path.read_bytes())
                for start in range(0, len(text), 6000):
                    self.segments.append({"artifact_id": artifact["id"], "char_start": start, "text": text[start:start+6000]})
            else:
                record["status"] = "unsupported"
                self.notes.append(artifact["name"] + "：该格式未解析")
        except InterruptedError:
            raise
        except Exception:
            record["status"] = "failed"
            self.notes.append(artifact["name"] + "：解析出现异常，请核对格式、加密及解析限额")

    def pdf(self, path, artifact, record):
        with PDF_LOCK:
            doc = pdfium.PdfDocument(str(path))
            try:
                count = len(doc)
                record.update(pages_total=count, pages_processed=min(count, MAX_PAGES))
                for index in range(min(count, MAX_PAGES)):
                    if self.cancelled():
                        raise InterruptedError()
                    page = doc[index]
                    try:
                        tp = page.get_textpage()
                        try:
                            text = tp.get_text_range()
                        finally:
                            tp.close()
                        self.segments.append({"artifact_id": artifact["id"], "page": index+1, "text": text})
                        if not text.strip():
                            self.notes.append(f"{artifact['name']} 第{index+1}页无文本层，OCR未启用")
                        w,h = page.get_size()
                        bitmap = page.render(scale=min(1.7, 1250/max(w,h)))
                        try:
                            rendered = bitmap.to_pil().convert("RGB").copy()
                        finally:
                            bitmap.close()
                        self.image(rendered, artifact, index+1, [0,0,1,1], "page")
                        nested = False
                        for obj in page.get_objects(max_depth=1):
                            if obj.type == pdfium.raw.FPDF_PAGEOBJ_FORM:
                                nested = True
                            if obj.type != pdfium.raw.FPDF_PAGEOBJ_IMAGE:
                                continue
                            box = pdf_box(page, obj.get_bounds(), rendered.width, rendered.height)
                            if box[2] <= box[0] or box[3] <= box[1]:
                                continue
                            crop = rendered.crop(tuple(round(v * (rendered.width if n%2==0 else rendered.height)) for n,v in enumerate(box)))
                            self.image(crop, artifact, index+1, box, "object")
                        if nested:
                            self.notes.append(f"{artifact['name']} 第{index+1}页包含Form对象，以整页渲染后备检查，未单独映射其内部图片")
                    finally:
                        page.close()
                if count > MAX_PAGES:
                    record["status"] = "partial"
                    self.notes.append(artifact["name"] + f"：仅解析前{MAX_PAGES}页")
            finally:
                doc.close()


def table_evidence_image(source, anchor, table_cache=None):
    key=str(Path(source).resolve())
    if table_cache is None:
        tables, _ = read_tables(source)
    else:
        if key not in table_cache:table_cache[key]=read_tables(source)[0]
        tables=table_cache[key]
    table = next((t for t in tables if t["sheet"] == anchor.get("sheet")), None)
    if table is None:
        raise ValueError("找不到指定工作表")
    r0,r1,c0,c1 = anchor["cells"]
    rows = table["rows"]
    if r1 > len(rows) or c1 > max((len(r) for r in rows), default=0):
        raise ValueError("单元格范围超出解析范围")
    if r1-r0 > 15 or c1-c0 > 7:
        raise ValueError("证据视图范围过大")
    ncols=c1-c0+1
    width, height = max(440,80+ncols*210), 85 + (r1-r0+2)*42
    image = Image.new("RGB", (width,height), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 18)
    except OSError:
        font = ImageFont.load_default()
    draw.text((20,8), "原始单元格视图（非屏幕截图）", font=font, fill="#24594e")
    draw.text((20,34), str(anchor['sheet'])[:35], font=font, fill="#24594e")
    cellw = (width-100)//ncols
    for ir, row_index in enumerate(range(r0-1,r1)):
        y = 65+ir*42
        draw.text((10,y+8), str(row_index+1), font=font, fill="#64756f")
        row = rows[row_index]
        for ic, column in enumerate(range(c0-1,c1)):
            x = 75+ic*cellw
            draw.rectangle((x,y,x+cellw,y+42), outline="#d4ded8")
            draw.text((x+6,y+8), str(row[column] if column<len(row) else "")[:18], font=font, fill="#203b36")
    return image
