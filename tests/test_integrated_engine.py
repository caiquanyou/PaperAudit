"""Integration calibration for migrated operations; no external/model calls."""
from __future__ import annotations

import json
import shutil
import sys
import unittest
import uuid
import zipfile
from unittest.mock import patch
from types import SimpleNamespace
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
import numpy as np
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
import openpyxl
from desktop_app.storage import Store
from desktop_app.parsers import Parser,read_tables
from desktop_app.rules import table_checks
from desktop_app.rules import text_checks
from desktop_app.image_checks import run_image_checks,merge_findings,Features,geometric_matches
from desktop_app.audit_engine import AuditEngine,DIMENSIONS
from desktop_app.models_client import ModelError,ModelClient,endpoint_url
from desktop_app.service import Service
from desktop_app.reports import generate_report,PDF_LOCK
import pypdfium2 as pdfium

WORK=ROOT/'.build/integrated-tests'


class IntegratedTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        WORK.mkdir(parents=True,exist_ok=True)
        rng=np.random.default_rng(42)
        cls.image=WORK/'synthetic-texture.png'
        Image.fromarray(rng.integers(0,256,(360,420,3),dtype=np.uint8)).save(cls.image)
        Image.new('RGB',(180,180),'#a7a7a7').save(WORK/'low.png')
        Image.new('RGB',(160,160),(0,129,0)).save(WORK/'green.png')
        cls.pdf=WORK/'synthetic-reuse.pdf'
        c=canvas.Canvas(str(cls.pdf),pagesize=(600,800))
        c.drawString(40,760,'SYNTHETIC CALIBRATION - NOT RESEARCH')
        c.drawImage(str(cls.image),40,360,width=240,height=200)
        c.drawImage(str(cls.image),320,360,width=240,height=200)
        c.drawString(40,300,'Same image deliberately drawn twice as a positive control.')
        c.showPage()
        c.drawString(40,760,'SECOND PAGE - tail methods and supplements must be included.')
        c.drawString(40,710,'p = 1.5 is a deliberately invalid statistical value.')
        c.save()
        cls.xlsx=WORK/'synthetic-supplement.xlsx'
        wb=openpyxl.Workbook();ws=wb.active;ws.title='Measurements'
        ws.append(['Sample','Measurement A','Measurement B','Display precision'])
        for i in range(1,36):
            ws.append([f'S{i}',i*.37+1.11,(i*.37+1.11)*2+3,.50 if i%2 else .55])
            ws.cell(i+1,4).number_format='0.00'
        ws.append(['Err','#REF!',2,.50]);ws.row_dimensions[4].hidden=True
        wb.save(cls.xlsx)

    def store(self):
        return Store(WORK/uuid.uuid4().hex[:10])

    def test_pdf_occurrences_keep_two_placements_and_page_coordinates(self):
        store=self.store();paper=store.import_paper(self.pdf)
        parser=Parser(store).parse(store.artifacts(paper))
        objects=[r for r in parser.images if r['kind']=='object']
        self.assertEqual(len(objects),2)
        self.assertNotEqual(objects[0]['bbox'],objects[1]['bbox'])
        self.assertTrue(all(0<=v<=1 for r in objects for v in r['bbox']))
        self.assertTrue(any('SECOND PAGE' in s['text'] for s in parser.segments))

    def test_legacy_numeric_operations_preserve_precision_and_locators(self):
        tables,notes=read_tables(self.xlsx)
        self.assertEqual(tables[0]['rows'][1][3],'0.50')
        for table in tables:table['artifact_id']='fixture'
        findings,checks=table_checks(tables)
        rules={f['rule_id'] for f in findings}
        self.assertTrue({'constant_increment','affine_columns','formula_error','hidden_regions'}<=rules)
        self.assertTrue(all(a.get('cells') for f in findings for a in f['anchors']))
        self.assertTrue(any(c['name']=='Benford' and c['status']=='needs_mapping' for c in checks))

    def test_archive_children_persist_without_changing_input_fingerprint(self):
        source=WORK/'supplement.zip'
        with zipfile.ZipFile(source,'w') as z:z.write(self.xlsx,'tables/measure.xlsx')
        store=self.store();paper=store.import_paper(self.pdf,[source]);fingerprint=store.fingerprint(paper)
        parsed=Parser(store).parse(store.artifacts(paper))
        children=[a for a in store.artifacts(paper) if a.get('parent_id')]
        self.assertEqual(len(children),1)
        self.assertEqual(store.fingerprint(paper),fingerprint)
        self.assertTrue(store.resolve(children[0]['relative_path']).is_file())
        self.assertEqual(parsed.tables[0]['artifact_id'],children[0]['id'])

    def test_zip_traversal_is_rejected(self):
        source=WORK/'unsafe.zip'
        with zipfile.ZipFile(source,'w') as z:z.writestr('../outside.txt','forbidden')
        store=self.store();paper=store.import_paper(self.pdf,[source])
        parsed=Parser(store).parse(store.artifacts(paper))
        self.assertTrue(any(c['status']=='failed' for c in parsed.coverage))
        self.assertFalse(any(a.get('parent_id') for a in store.artifacts(paper)))

    def test_rotated_cropped_pdf_maps_image_regions(self):
        from pypdf import PdfReader,PdfWriter
        source=PdfReader(str(self.pdf));page=source.pages[0]
        page.cropbox.lower_left=(20,100);page.cropbox.upper_right=(580,740);page.rotate(90)
        writer=PdfWriter();writer.add_page(page)
        path=WORK/'rotated-cropped.pdf'
        with path.open('wb') as handle:writer.write(handle)
        store=self.store();paper=store.import_paper(path)
        parsed=Parser(store).parse(store.artifacts(paper))
        images=[r for r in parsed.images if r['kind']=='object']
        self.assertEqual(len(images),2)
        for record in images:
            with Image.open(record['path']) as im:
                self.assertGreater(np.array(im).std(),30)
                self.assertAlmostEqual(im.width/im.height,200/240,delta=.04)

    def test_horizontal_increment_and_identifier_conflict_are_located(self):
        table={'artifact_id':'test','sheet':'Measurements','rows':[['Sample','A','B','C','D','E','F'],['S1',1.3,2.6,3.9,5.2,6.5,7.8],['S1',2.3,3.6,4.9,6.2,7.5,8.8]]}
        findings,_=table_checks([table])
        self.assertTrue({'horizontal_increment','identifier_mapping'}<={f['rule_id'] for f in findings})

    def image_record(self,path,index):
        return {'id':str(index),'artifact_id':'fixture-'+str(index),'page':1,'bbox':[0,0,1,1],'kind':'raster','path':str(path)}

    def test_exact_low_texture_and_flip_crop_channels(self):
        low=WORK/'low.png';low2=WORK/'low2.png'
        Image.new('RGB',(180,180),'#a7a7a7').save(low);shutil.copyfile(low,low2)
        findings,_=run_image_checks([self.image_record(low,1),self.image_record(low2,2)])
        self.assertTrue(any(f['metrics'].get('exact_pixels') for f in findings))
        flipped=WORK/'flipped-crop.png'
        with Image.open(self.image) as im:im.crop((30,30,360,300)).transpose(Image.Transpose.FLIP_LEFT_RIGHT).save(flipped)
        records=[self.image_record(self.image,3),self.image_record(flipped,4)]
        features=Features(records)
        matches=geometric_matches(features,*records,mode='hflip')
        self.assertTrue(matches,'explicit mirror and crop should pass geometry/pixel calibration')
        self.assertGreater(matches[0]['metrics']['pixel_correlation'],.85)

    def test_duplicate_algorithm_findings_merge_by_source_regions(self):
        pair=[{'artifact_id':'a','page':1,'bbox':[.1,.1,.5,.5]},{'artifact_id':'b','page':2,'bbox':[.2,.2,.6,.6]}]
        values=[{'kind':'image','anchors':pair,'methods':[method],'observation':method} for method in ['pHash','ORB','SIFT']]
        result=merge_findings(values)
        self.assertEqual(len(result),1)
        self.assertEqual(len(result[0]['methods']),3)

    def test_same_visible_pdf_region_is_not_image_reuse(self):
        left={'artifact_id':'a','page':4,'bbox':[.2,.12,.86,.8],'occurrence_id':'first-layer'}
        right={**left,'occurrence_id':'second-layer'}
        self.assertEqual(merge_findings([{'kind':'image','anchors':[left,right]}]),[])
        different={**right,'page':5}
        self.assertEqual(len(merge_findings([{'kind':'image','anchors':[left,different]}])),1)

    def test_local_engine_generates_real_evidence_reports_without_key(self):
        store=self.store();paper=store.import_paper(self.pdf,[self.xlsx])
        service=Service(store,AuditEngine(store))
        job=service.enqueue(paper);service.run_once()
        self.assertEqual(store.job(job)['status'],'partial')
        result=json.loads(store.result_path(store.job(job)).read_text(encoding='utf-8'))
        self.assertIn('0.6.2',result['engine_id'])
        self.assertTrue(any(f['kind']=='image' for f in result['findings']))
        self.assertTrue(any(f['rule_id']=='p_value_range' for f in result['findings']))
        folder=store.resolve(store.report_list(paper)[0]['relative_dir'])
        md=(folder/'report.md').read_text(encoding='utf-8')
        self.assertIn('表格单元格视图',md)
        evidence=json.loads((folder/'evidence.json').read_text(encoding='utf-8'))
        self.assertTrue(all(e['status']=='captured' for group in evidence['evidence'] for e in group))
        output=ROOT/'output/整合引擎_合成校准报告'
        shutil.copytree(folder,output,dirs_exist_ok=True)
        with PDF_LOCK:
            doc=pdfium.PdfDocument(str(folder/'report.pdf'))
            for index in range(len(doc)):
                page=doc[index];bitmap=page.render(scale=.9)
                bitmap.to_pil().save(WORK/f'calibration-report-{index+1}.png')
                bitmap.close();page.close()
            print('Calibration report pages:',len(doc))
            doc.close()

    def test_model_failure_never_becomes_low_risk_and_local_output_survives(self):
        class BadClient:
            def __init__(self,*_):pass
            def chat(self,*_):raise ModelError('http_401')
        store=self.store();paper=store.import_paper(self.pdf)
        engine=AuditEngine(store,client_factory=BadClient)
        result=engine(store.paper(paper),store.artifacts(paper),{'base_url':'https://example.invalid','text_model':'chosen-model'},lambda:'TEST-ONLY')
        self.assertTrue(any(c['status']=='failed' and c['name'].startswith('模型六维') for c in result['checks']))
        self.assertTrue(result['partial'])
        self.assertTrue(result['findings'])
        self.assertNotIn('risk_level',result)

    def test_model_tail_coverage_and_unverified_quote_rejection(self):
        received=[]
        class Client:
            def __init__(self,*_):pass
            def chat(self,system,text,*_):
                received.append(text)
                return {'dimensions':{k:'insufficient_data' for k in DIMENSIONS},'findings':[{'title':'fabricated quote','evidence':[{'segment_id':'s00000_0','quote':'not present in this paper'}]}]}, {'total_tokens':42}
        long=WORK/'long-supplement.txt'
        long.write_text(('Repeated synthetic methods content, not research data. '*1000)+' FINAL_TAIL_MARKER_50000_CHARS',encoding='utf-8')
        store=self.store();paper=store.import_paper(self.pdf,[long])
        engine=AuditEngine(store,client_factory=Client)
        result=engine(store.paper(paper),store.artifacts(paper),{'base_url':'https://example.invalid','text_model':'chosen-model'},lambda:'TEST-ONLY')
        self.assertTrue(any('SECOND PAGE' in t for t in received))
        self.assertTrue(any('FINAL_TAIL_MARKER_50000_CHARS' in t for t in received))
        self.assertFalse(any(f['title']=='fabricated quote' for f in result['findings']))
        self.assertTrue(any('未通过原文引文定位' in n for n in result['limitations']))

    def test_provider_address_policy(self):
        self.assertEqual(endpoint_url('http://127.0.0.1:8000/v1',True),'http://127.0.0.1:8000/v1')
        for url in ['http://outside.invalid','https://user:pass@example.invalid','https://example.invalid?key=value']:
            with self.assertRaises(ValueError):endpoint_url(url,True)

    def test_deepseek_structured_mode_and_failed_usage_preserved(self):
        client=ModelClient({'base_url':'https://api.deepseek.com'},'TEST-NOT-REAL')
        sent=[]
        def response(path,payload):
            sent.append(payload)
            return {'choices':[{'finish_reason':'length','message':{'content':''}}],'usage':{'prompt_tokens':12,'completion_tokens':4096,'total_tokens':4108}}
        with patch.object(client,'request',side_effect=response):
            with self.assertRaises(ModelError) as error:client.chat('JSON','test','deepseek-flash')
        self.assertEqual(sent[0]['thinking'],{'type':'disabled'})
        self.assertEqual(sent[0]['response_format'],{'type':'json_object'})
        self.assertEqual(error.exception.usage['total_tokens'],4108)
        other=ModelClient({'base_url':'https://other.invalid/v1'},'TEST-NOT-REAL')
        self.assertNotIn('thinking',other.generation_options())

    def test_scientific_p_values_are_not_truncated_into_false_alarms(self):
        text='p=5e-3; p=5E−3; p=5×10⁻³; p=5 x 10^-3; p=.005; p=1.5.'
        findings,_=text_checks([{'artifact_id':'test','page':1,'text':text}])
        self.assertEqual(len(findings),1)
        self.assertIn('1.5',findings[0]['observation'])
        findings,checks=text_checks([{'artifact_id':'test','page':1,'text':'p=5e- missing exponent'}])
        self.assertEqual(findings,[])
        self.assertTrue(any(c['name']=='p值指数表达解析' for c in checks))

    def test_equal_grayscale_does_not_mean_equal_rgb_pixels(self):
        a,b=WORK/'red.png',WORK/'green.png'
        Image.new('RGB',(160,160),(255,0,0)).save(a)
        Image.new('RGB',(160,160),(0,129,0)).save(b)
        findings,_=run_image_checks([self.image_record(a,81),self.image_record(b,82)])
        self.assertFalse(any(f['metrics'].get('exact_pixels') for f in findings))

    def test_page_self_copy_runs_even_when_normal_images_exist(self):
        a=self.image_record(self.image,91);b={**self.image_record(self.image,92),'kind':'page'}
        visited=[]
        with patch('desktop_app.image_checks.self_copy',side_effect=lambda f,r:visited.append(r['id'])):
            run_image_checks([a,b])
        self.assertEqual(visited,['91','92'])

    def test_external_partial_is_terminal_and_remains_partial(self):
        store=self.store();paper=store.import_paper(self.pdf);service=Service(store)
        job=service.enqueue(paper);service.run_once()
        result={'job_id':job,'input_fingerprint':store.job(job)['input_fingerprint'],'engine_id':'external-test-only','findings':[],'checks':[{'name':'scope','status':'unavailable'}],'partial':True}
        service.accept_external_result(paper,result)
        self.assertEqual(store.job(job)['status'],'partial')
        with self.assertRaises(ValueError):service.accept_external_result(paper,result)

    def test_vision_cache_binds_images_and_source_locations(self):
        requests=[]
        class Client:
            def __init__(self,*_):pass
            def chat(self,system,text,model,images=()):
                requests.append((model,tuple(images)))
                return {'summary':'test-only visual observation','dimensions':{k:'insufficient_data' for k in DIMENSIONS},'findings':[]},{'total_tokens':1}
        store=self.store();paper=store.import_paper(self.pdf)
        records=[self.image_record(self.image,1),self.image_record(WORK/'low.png',2),self.image_record(WORK/'green.png',3)]
        candidates=[{'kind':'image','observation':'identical candidate description','anchors':[{'artifact_id':'f','page':1,'bbox':[.1,.1,.2,.2],'occurrence_id':str(i)}]} for i in [1,2,3]]
        engine=AuditEngine(store,client_factory=Client);calls=[];checks=[];notes=[]
        engine.model_review(store.paper(paper),{'base_url':'https://example.invalid','text_model':'text-test','vision_model':'vision-test','max_model_calls':8},'TEST-ONLY',{'s0':{'artifact_id':'f','page':1,'text':'A synthetic snippet.'}},candidates,checks,notes,calls,SimpleNamespace(images=records))
        visual=[c for c in calls if c['image_count']]
        self.assertEqual(len(visual),3)
        self.assertEqual(len({c['input_hash'] for c in visual}),3)
        self.assertTrue(all(c.get('image_hashes') and not c.get('cached') for c in visual))


if __name__=='__main__':unittest.main(verbosity=2)
